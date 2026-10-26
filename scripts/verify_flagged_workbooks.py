"""Read-only verification of the narrow XLSX patches and preserved evidence."""
import json
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'outputs/flagged-outlets-20261026'
BACK=ROOT/'recovery/flagged-outlets-20261026';STAGE=OUT/'staged'
patches=json.loads((OUT/'workbook_patches.json').read_text())
checks={}
empty=lambda x: None if x=='' else x
for name,p in patches.items():
 old=load_workbook(BACK/name,data_only=False);new=load_workbook(STAGE/name,data_only=False)
 assert old.sheetnames==new.sheetnames,(name,'sheets')
 expected={}
 if p.get('sheet'):expected[p['sheet']]={(r,c):v for r,c,v in p['edits']}
 else:expected={s:{(r,c):v for r,c,v in v['edits']} for s,v in p['sheets'].items()}
 formulas=0
 for s in old:
  t=new[s.title];changed=expected.get(s.title,{})
  assert list(s.merged_cells.ranges)==list(t.merged_cells.ranges),(name,s.title,'merges')
  for row in s:
   for cell in row:
    wanted=changed.get((cell.row,cell.column),cell.value)
    got=t.cell(cell.row,cell.column).value
    assert empty(got)==empty(wanted),(name,s.title,cell.coordinate,wanted,got)
    if cell.data_type=='f':formulas+=1
  for (r,c),v in changed.items():assert empty(t.cell(r,c).value)==empty(v),(name,s.title,r,c)
 if p.get('append'):
  t=new[p['sheet']]
  assert t.max_row==p['original_rows']+len(p['append'])
  for i,row in enumerate(p['append'],p['original_rows']+1):
   assert [empty(t.cell(i,c+1).value) for c in range(len(row))]==[empty(v) for v in row],(name,i)
 checks[name]={'original_formulas_preserved':formulas,'unrelated_values_preserved':True,'added_rows':len(p.get('append',[]))}
bridge=load_workbook(STAGE/'bridge_outlets.xlsx',data_only=True).active
data=json.loads((STAGE/'bridge.json').read_text())
assert list(bridge.values)==[('outlet_final','total_attendance')]+[(r.get('outlet_final'),r['total_attendance']) for r in data]
(OUT/'workbook_validation.json').write_text(json.dumps(checks,indent=2)+'\n')
print(checks);print('Bridge verified against staged physical attendance counts.')
