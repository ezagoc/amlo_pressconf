"""Prepare narrow workbook patches from reviewed evidence; never author XLSX here."""
import json,csv,re,unicodedata
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'outputs/flagged-outlets-20261026';BACK=ROOT/'recovery/flagged-outlets-20261026'
inv=json.loads((OUT/'initial_inventory.json').read_text());cfg=json.loads((OUT/'decisions.json').read_text())
norm=lambda s:re.sub('[^a-z0-9]','',unicodedata.normalize('NFKD',str(s or '').lower()).encode('ascii','ignore').decode())
with open('/tmp/sheinbaum-current-mapped.csv') as f:q=list(csv.DictReader(f))
groupmap={b:g['canonical'] for g in cfg['groups'] for b in g['brands']}
meta_keys={norm(name):g['canonical'] for g in cfg['groups'] for name in g['metadata_names']}
meta_keys.update({norm(k):v for k,v in {'Formula':'grupo formula','Noticieros el Reloj de la Tlaxiaqueña':'noticias el reloj de la tlaxiaquena','Canal Universidad de Guadalajara':'universidad de guadalajara'}.items()})
def canonical(o):return groupmap.get(o,{'capital':'capital puebla'}.get(o,o))
display={'el heraldo':'El Heraldo de México','noticieros hctv':'Noticieros HCTV','telesur':'Telesur','alcanzando el conocimiento':'Alcanzando el Conocimiento','canal del congreso':'Canal del Congreso','zagazine':'Zagazine','subjetivo mx':'Subjetivo MX','10 am mx':'10 AM MX','radio educacion':'Radio Educación','revista zocalo':'Revista Zócalo','canal 22':'Canal 22','enlace 13 television':'Enlace 13 Televisión','poder mexico':'Poder México','dime noticias':'DIME Noticias','canal 5 tv cozumel':'Canal 5 TV Cozumel','global energy':'Global Energy','comunicante mx':'Comunicante MX','afp':'Agence France-Presse (AFP)','alerta multimedia':'Alerta Multimedia','infopolitano':'Infopolitano','le monde':'Le Monde','lobos cadena 7':'Lobos Cadena 7','pulso sur':'Pulso Sur','revolucion 30':'Revolución 3.0','ruido en la red':'Ruido en la Red','siker':'Siker','telemundo':'Telemundo','chilango':'Chilango','noti 13':'Noti 13','te lo cuento':'Te lo Cuento','al tanto guerrero':'Al Tanto Guerrero','el ceo':'El CEO','la izquierda noticias':'La Izquierda Noticias','la orquesta mx':'La Orquesta MX','la silla rota':'La Silla Rota','neurona digital noticias':'Neurona Digital Noticias','nopal times':'Nopal Times','quequi':'Quequi','somos mas':'Somos Más','tv baja':'TV Baja','yo tambien':'Yo También','bctneus':'BCTneus','capital':'Capital Puebla','cnr canal 73':'CNR Canal 73','el momento':'El Momento','el portal del noreste':'El Portal del Noreste','enprivadoveracruzcommx':'En Privado Veracruz','fusion latina chicago':'Fusión Latina Chicago','nbc news':'NBC News','news san miguel':'News San Miguel','noticias el agro':'Noticias El Agro','periodistas del sur':'Periodistas del Sur','pridemx':'PrideMX','radio hit la xplosiva':'Radio Hit La Xplosiva','radio tosepan limakxtum':'Radio Tosepan Limakxtum','sandraaguileracom':'SandraAguilera.com','cultura colectiva':'Cultura Colectiva','reyna haydee (independiente)':'Reyna Haydeé (independiente)','associated press':'Associated Press','el regreso de donald trump':'El Regreso de Donald Trump','agora digital':'Ágora Digital','lja':'LJA','tv migrante':'TV Migrante','el fildeo':'El Fildeo','efecto colateral':'Efecto Colateral','radio resultados':'Radio Resultados','domo de cristal':'Domo de Cristal','ap noticias':'AP Noticias','agora celaya':'Ágora (Celaya)'}
display.update({'bloomberg linea':'Bloomberg Línea','bloomberg news':'Bloomberg News','bloomberg':'Bloomberg (unspecified service)'})
display['abc radio']='ABC Radio'
missing=[r['outlet'] for r in inv if not r['metadata_names']]+['ap noticias','agora celaya','bloomberg linea','bloomberg news','bloomberg']
if 'abc radio' in groupmap:missing.append('abc radio')
# Keep distinct brands even when their conference ID has been grouped.
patches={};additions={}
for filename,sheet in [('mexican_newspapers.xlsx','Periodicos'),('newspaper_metadata.xlsx','Hoja1')]:
 w=load_workbook(BACK/filename,data_only=False);s=w[sheet];headers=[c.value for c in s[1]];cols={h:i+1 for i,h in enumerate(headers) if h};vals=list(s.values)
 existing={norm(r[cols['name.page']-1]) for r in vals[1:] if r[cols['name.page']-1]}
 edits=[]
 for rownum,row in enumerate(vals[1:],2):
  name=row[cols['name.page']-1];target=meta_keys.get(norm(name))
  if not target:continue
  g=next((g for g in cfg['groups'] if g['canonical']==target),None)
  if filename.startswith('mexican'):
   for key in ['names.conference','names.transcripts']:edits.append([rownum,cols[key],target])
  else:
   old=row[cols['verification.reason']-1] or ''
   note=g['basis']+' Preserve separate brands, ownership fields and channels.' if g else 'Reviewed 2026-10-26: standardized conference source key '+target+'; exact source-label variants recorded in the reviewed alias registry.'
   edits.append([rownum,cols['verification.reason'],old+'\n'+note])
  if g:edits.append([rownum,cols['outlet_aliases'],'; '.join(g['brands'])+' (shared conference records; separate brands)'])
 new=[]
 # Later additions go last to preserve the metadata IDs of earlier new brands.
 for outlet in sorted(set(missing),key=lambda o:(o=='abc radio',o)):
  name=display.get(outlet,outlet.title())
  if norm(name) in existing:continue
  rows=[r for r in q if r.get('outlet_brand_aux')==outlet or r['outlet_aux']==outlet]
  if outlet=='capital':rows=[r for r in q if r['outlet_aux_before']=='capital' and r['reporter_aux']=='abel cuapa']
  if outlet=='agora digital':rows=[r for r in rows if r['reporter_aux']!='gabriela montejano']
  if outlet=='associated press':rows=[r for r in rows if r['reporter_aux']!='alejandro perez']
  if outlet=='abc radio':rows=[r for r in q if r.get('outlet_brand_aux')=='asi lo dice la mont']
  if not rows:continue
  target=canonical(outlet);reporters=sorted({r['reporter'] for r in rows if r['reporter'] not in ['','NA']})
  evidence=next((r for r in rows if r['outlet'] not in ['','NA']),rows[0]);source=evidence['source_url']
  note='Identity review 2026-10-26: '+name+' is identified in official transcripts. Conference key: '+target+'. Source: '+source
  if outlet=='abc radio':note='Identity review 2026-10-26: ABC RADIO 760 AM occurs in the raw attendance roster on all six covered Federico La Mont asking dates. Historical program affiliation documented at https://gaceta.diputados.gob.mx/PDF/InfoDip/64/1052-20210817-I-II-III.pdf . Conference key: audiorama. Do not confuse with ABC de Monterrey, ABC Australia or unspecified ABC News.'
  g=next((g for g in cfg['groups'] if outlet in g['brands']),None)
  if g:note+=' '+g['basis']+' Retain this brand separately.'
  note+=' Legal identity, ownership, founding year and address have not been established in this identity review.'
  if outlet in cfg['notes']:note+=' '+cfg['notes'][outlet]
  row=[None]*s.max_column
  def put(k,v):row[cols[k]-1]=v
  put('name.page',name)
  if filename.startswith('mexican'):
   put('names.transcripts',target);put('names.conference',target);put('journalist.mañanera','; '.join(reporters))
  else:
   put('needs.verification','yes');put('verification.reason',note)
  if g:put('outlet_aliases','; '.join(g['brands'])+' (shared conference records; separate brands)')
  if outlet=='capital':put('url','https://capitalpuebla.com/')
  if outlet=='fusion latina chicago':put('url','https://www.fusionlatinachicago.org/')
  if outlet=='noticias el agro':put('url','https://elagro.com.mx/')
  if outlet=='radio tosepan limakxtum':put('url','https://radio.tosepan.coop/')
  new.append(row);existing.add(norm(name))
 patches[filename]={'sheet':sheet,'original_rows':s.max_row,'original_cols':s.max_column,'edits':edits,'append':new}
 additions[filename]=[r[cols['name.page']-1] for r in new]
(OUT/'workbook_patches.json').write_text(json.dumps(patches,ensure_ascii=False,indent=2)+'\n')
(OUT/'metadata_additions.json').write_text(json.dumps(additions,ensure_ascii=False,indent=2)+'\n')
print({f:{'edits':len(v['edits']),'new_brand_rows':len(v['append'])} for f,v in patches.items()})
