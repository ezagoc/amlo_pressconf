from pathlib import Path
import json
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
out=root/'outputs/teleurban-20261026'
back=root/'recovery/teleurban-20261026'
for name,sheet in [('mexican_newspapers.xlsx','Periodicos'),('newspaper_metadata.xlsx','Hoja1')]:
    before=load_workbook(back/name,data_only=False)
    after=load_workbook(out/'staged'/name,data_only=False)
    changes={(r,c):v for r,c,v in json.loads((out/(name+'.changes.json')).read_text())}
    assert before.sheetnames==after.sheetnames
    count=0
    for s in before.sheetnames:
        a,b=before[s],after[s]
        assert a.max_row==b.max_row and a.max_column==b.max_column
        for row in a:
            for cell in row:
                expected=changes.get((cell.row,cell.column),cell.value) if s==sheet else cell.value
                actual=b.cell(cell.row,cell.column).value
                assert actual==expected,(name,s,cell.coordinate,expected,actual)
                count+=cell.data_type=='f'
    print(name,': only approved cell changes; preserved formulas:',count)
